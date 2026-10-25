"""任务内启动程序。业务权限、联网限制与容器期限由可信管理侧执行。"""
import argparse
import base64
import ast
import csv
from contextlib import closing
import ipaddress
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import selectors
import signal
import stat
import subprocess
import sys
import tempfile
import time
from urllib.parse import urlsplit
import zipfile

try:
    from .export import MAX_FILE_BYTES, MAX_TOTAL_BYTES, read_artifacts
except ImportError:
    from export import MAX_FILE_BYTES, MAX_TOTAL_BYTES, read_artifacts

MAX_INPUT_BYTES = 8 * 1024 * 1024
MAX_LOG_BYTES = 64 * 1024
MAX_TEXT = 200000
MAX_ROWS = 2000
MAX_WEB_BYTES = 4 * 1024 * 1024
MAX_IMAGE_PIXELS = 16_000_000


def bounded_file(path: Path, limit=MAX_INPUT_BYTES) -> bytes:
    # 多读取一个字节来判断超限，避免先把整个文件装入内存。
    with path.open("rb") as stream:
        content = stream.read(limit + 1)
    if len(content) > limit:
        raise ValueError(f"输入超过限制：{path.name}")
    return content


def task_path(workspace: Path, relative: str, base="") -> Path:
    # 同时检查字面路径和解析后的路径，避免相对路径或目录链接越界。
    value = PurePosixPath(relative)
    if not relative or value.is_absolute() or ".." in value.parts or "\\" in relative:
        raise ValueError("任务路径必须在工作目录内")
    root = (workspace / base).resolve()
    result = root / relative
    if not result.resolve().is_relative_to(root):
        raise ValueError("任务路径越出工作目录")
    return result


def input_file(task, workspace, suffixes=None):
    # 多个候选文件时要求调用方明确指定，防止悄悄解析错误输入。
    chosen = task.get("options", {}).get("input")
    if chosen:
        path = task_path(workspace, chosen, "inputs")
    else:
        candidates = sorted(path for path in (workspace / "inputs").iterdir()
                            if path.is_file() and (suffixes is None or path.suffix.lower() in suffixes))
        if len(candidates) != 1:
            raise ValueError("请通过 options.input 指定一个输入文件")
        path = candidates[0]
    if path.is_symlink() or not path.is_file():
        raise ValueError("输入必须是普通文件")
    bounded_file(path)
    return path


def write_output(out: Path, name: str, content: bytes):
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("单个输出超过 2 MiB")
    descriptor = os.open(out / name, os.O_WRONLY | os.O_CREAT | os.O_TRUNC | os.O_NOFOLLOW, 0o600)
    with os.fdopen(descriptor, "wb") as stream:
        stream.write(content)


def write_output_json(out, name, content):
    write_output(out, name, (json.dumps(content, ensure_ascii=False, indent=2, default=str) + "\n").encode())


def execute_script(task, workspace):
    """只执行调用方指定的入口脚本；Skill 步骤规划由沙箱外的 Agent 决定。"""
    entrypoint = task_path(workspace, task.get("entrypoint", ""))
    if not entrypoint.is_file() or entrypoint.is_symlink():
        raise ValueError("找不到执行脚本")
    interpreter = "node" if task.get("environment") == "node" else sys.executable
    if interpreter == "node" and entrypoint.suffix not in (".js", ".mjs", ".cjs"):
        raise ValueError("Node 环境需要 JavaScript 脚本")
    if interpreter != "node" and entrypoint.suffix != ".py":
        raise ValueError("Python 环境需要 .py 脚本")
    # 动态依赖只影响任务解释器；平台控制程序继续使用基础 Python。
    if interpreter != 'node' and Path('/opt/env/python/bin/python').is_file():
        interpreter = '/opt/env/python/bin/python'
    if interpreter == 'node' and Path('/opt/env/node/node_modules').is_dir():
        (workspace / 'node_modules').symlink_to('/opt/env/node/node_modules', target_is_directory=True)
    process = subprocess.Popen([interpreter, str(entrypoint)], cwd=workspace,
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=True)
    selector = selectors.DefaultSelector()
    # 同时排空两路管道，避免子进程因某一路日志管道填满而阻塞。
    logs = {"stdout": bytearray(), "stderr": bytearray()}
    selector.register(process.stdout, selectors.EVENT_READ, "stdout")
    selector.register(process.stderr, selectors.EVENT_READ, "stderr")
    deadline = time.monotonic() + 120
    try:
        while selector.get_map():
            if time.monotonic() > deadline:
                raise TimeoutError("脚本超过启动程序的 120 秒上限")
            for key, _ in selector.select(timeout=0.2):
                chunk = os.read(key.fileobj.fileno(), 16384)
                if not chunk:
                    selector.unregister(key.fileobj)
                    continue
                if len(logs[key.data]) + len(chunk) > MAX_LOG_BYTES:
                    raise ValueError("脚本 stdout/stderr 超过 64 KiB")
                logs[key.data].extend(chunk)
        code = process.wait(timeout=1)
        decoded = {name: data.decode("utf-8", errors="replace").encode("utf-8")[:MAX_LOG_BYTES].decode("utf-8", errors="ignore")
                   for name, data in logs.items()}
        if code:
            raise RuntimeError(f"脚本退出码：{code}；{decoded['stderr'][-2000:]}")
        return decoded
    except Exception as error:
        error.logs = {name: data.decode("utf-8", errors="replace").encode("utf-8")[:MAX_LOG_BYTES].decode("utf-8", errors="ignore")
                      for name, data in logs.items()}
        raise
    finally:
        # 独立进程组允许连同脚本派生的子进程一起收尾。
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        process.wait()
        selector.close()
        process.stdout.close()
        process.stderr.close()


def safe_url(value):
    # 这里是 URL 的前置检查；域名解析及真实出网限制仍由外部代理执行。
    parsed = urlsplit(value)
    if (len(value) > 2048 or parsed.scheme not in ("http", "https") or not parsed.hostname
            or parsed.username or parsed.password or parsed.port not in (None, 80, 443)):
        raise ValueError("网页地址只支持无登录信息的 HTTP/HTTPS 公网 URL")
    try:
        address = ipaddress.ip_address(parsed.hostname)
    except ValueError:
        address = None
    if address is not None and not address.is_global:
        raise ValueError("网页地址不能是私网或保留地址")
    return value


def browser_task(task, workspace, out, render=False):
    """网页读取走受控代理，本地 HTML 渲染禁止外部资源请求。"""
    from playwright.sync_api import sync_playwright
    options = task.get("options", {})
    proxy = os.environ.get("HTTPS_PROXY") or os.environ.get("HTTP_PROXY")
    with tempfile.TemporaryDirectory(prefix="browser-", dir="/tmp") as profile, sync_playwright() as playwright:
        # 代理核验明文 SNI 并拒绝 ECH；新版 Chromium 已移除旧 feature flag。
        (Path(profile) / "Local State").write_text(json.dumps({"ssl": {"ech_enabled": False}}))
        launch = {"headless": True, "args": ["--no-sandbox", "--disable-dev-shm-usage"]}
        test_spki = os.environ.get("LAZYX_TEST_TLS_SPKI")
        if test_spki:
            if len(base64.b64decode(test_spki, validate=True)) != 32:
                raise ValueError("合成测试证书 SPKI 摘要不合法")
            launch["args"].append("--ignore-certificate-errors-spki-list=" + test_spki)
        if os.environ.get("PLAYWRIGHT_CHROMIUM_EXECUTABLE"):
            launch["executable_path"] = os.environ["PLAYWRIGHT_CHROMIUM_EXECUTABLE"]
        if proxy and not render:
            launch["proxy"] = {"server": proxy}
        width = min(max(int(options.get("width", 800)), 320), 1600)
        context = playwright.chromium.launch_persistent_context(profile, **launch,
            accept_downloads=False, java_script_enabled=not render,
            service_workers="block", viewport={"width": width, "height": 900})
        page = context.new_page()
        page.set_default_timeout(15000)
        if render:
            # 本地渲染关闭 JavaScript 并阻断资源加载，输出只依据上传内容。
            context.route("**/*", lambda route: route.abort())
            path = input_file(task, workspace, {".html", ".txt"})
            content = bounded_file(path).decode("utf-8")
            if path.suffix.lower() == ".txt":
                import html
                content = f"<html><meta charset='utf-8'><body><pre>{html.escape(content)}</pre></body></html>"
            page.set_content(content, wait_until="load")
            height = page.evaluate("Math.max(document.body.scrollHeight, document.documentElement.scrollHeight)")
            if height > 12000:
                raise ValueError("长图内容高度超过 12000 像素")
            format_name = task.get("render_format", "png")
            if format_name == "png":
                data = page.screenshot(full_page=True)
                write_output(out, "render.png", data)
            elif format_name == "pdf":
                data = page.pdf(format="A4", print_background=True)
                write_output(out, "render.pdf", data)
            else:
                raise ValueError("render_format 仅支持 png/pdf")
        else:
            # 浏览器下载也计入任务限额；覆盖正文、图片及脚本，不只检查最终文字。
            received = 0
            exceeded = False
            session = context.new_cdp_session(page)
            session.send("Network.enable")
            def limit_download(event):
                nonlocal received, exceeded
                received += max(event.get("dataLength", 0), event.get("encodedDataLength", 0))
                if received > MAX_WEB_BYTES and not exceeded:
                    exceeded = True
                    page.close()
            session.on("Network.dataReceived", limit_download)
            context.on("page", lambda popup: popup.close())
            def route_request(route):
                try:
                    safe_url(route.request.url)
                except ValueError:
                    route.abort()
                else:
                    route.continue_()
            context.route("**/*", route_request)
            try:
                response = page.goto(safe_url(task.get("url", "")), wait_until="load")
                safe_url(page.url)
                if response is not None and response.status in (401, 403):
                    raise ValueError("网页需要登录或没有访问权限")
                if response is None or response.status >= 400:
                    raise RuntimeError("网页请求失败")
                if page.locator('input[type="password"]').count():
                    raise ValueError("网页含登录表单，不处理账号登录")
                text = page.locator("body").inner_text()
                if len(text) > MAX_TEXT:
                    raise ValueError("网页正文超过长度限制")
                links = page.locator("a[href]").evaluate_all("nodes => nodes.slice(0,100).map(n => ({text:n.innerText.slice(0,200),url:n.href}))")
                write_output_json(out, "webpage.json", {"url": page.url, "title": page.title(), "text": text, "links": links})
            finally:
                # 达到限额后主动关闭页面；后续浏览器错误仍归为下载超量。
                if exceeded:
                    raise ValueError("网页下载超过 4 MiB") from None
        context.close()


def check_zip(source: Path, limit=MAX_TOTAL_BYTES):
    # 读取包结构前约束解压量、压缩比与路径，DOCX/XLSX 也复用此检查。
    with zipfile.ZipFile(source) as archive:
        entries = archive.infolist()
        if len(entries) > 100:
            raise ValueError("ZIP 文件项超过 100 个")
        total = 0
        for entry in entries:
            name = PurePosixPath(entry.filename)
            mode = entry.external_attr >> 16
            if (name.is_absolute() or ".." in name.parts or "\\" in entry.filename
                    or ":" in entry.filename or stat.S_ISLNK(mode) or entry.flag_bits & 1):
                raise ValueError("ZIP 含越界路径、符号链接或加密文件")
            total += entry.file_size
            if total > limit or entry.file_size > MAX_FILE_BYTES or entry.file_size / max(entry.compress_size, 1) > 100:
                raise ValueError("ZIP 解压大小或压缩比超过限制")
        return entries


def detected_type(source):
    """二进制依据签名及包结构；文本格式没有可靠魔数，仍需实际解析。"""
    data = bounded_file(source)
    declared = source.suffix.lower().lstrip('.')
    declared = {'jpg': 'jpeg'}.get(declared, declared)
    if data.startswith(b'%PDF-'):
        actual = 'pdf'
    elif data.startswith(b'\x89PNG\r\n\x1a\n'):
        actual = 'png'
    elif data.startswith(b'\xff\xd8\xff'):
        actual = 'jpeg'
    elif data.startswith(b'RIFF') and data[8:12] == b'WEBP':
        actual = 'webp'
    elif data.startswith((b'PK\x03\x04', b'PK\x05\x06')):
        entries = check_zip(source, 32 * 1024 * 1024)
        names = {entry.filename for entry in entries}
        word = {'[Content_Types].xml', 'word/document.xml'} <= names
        excel = {'[Content_Types].xml', 'xl/workbook.xml'} <= names
        if word and excel:
            raise ValueError('文件类型不明确：同时包含 Word 和 Excel 结构')
        actual = 'docx' if word else 'xlsx' if excel else 'zip'
    else:
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValueError('不支持的文件类型或文件头损坏') from None
        if any(ord(c) < 32 and c not in '\t\r\n' for c in text):
            raise ValueError('文本包含二进制控制字符')
        actual = declared if declared in ('txt', 'csv', 'json') else 'txt'
        if declared not in ('txt', 'csv', 'json') and text.lstrip().startswith(('{', '[')):
            json.loads(text)
            actual = 'json'
    if declared and declared != actual:
        raise ValueError(f'文件类型不符：扩展名 {declared}，内容识别为 {actual}')
    return actual


def png_bytes(image):
    """只输出像素，丢弃输入 EXIF、ICC 和文本元数据。"""
    from PIL import Image
    mode = 'RGBA' if 'A' in image.getbands() or 'transparency' in image.info else 'RGB'
    converted = image.convert(mode)
    with Image.frombytes(mode, converted.size, converted.tobytes()) as clean:
        buffer = io.BytesIO()
        clean.save(buffer, format='PNG')
    converted.close()
    data = buffer.getvalue()
    if len(data) > MAX_FILE_BYTES:
        raise ValueError('转换后的单张 PNG 超过 2 MiB，请降低 max_long_edge')
    return data


def normalize_image(source, actual, edge):
    """校正 EXIF 方向并等比缩小；输出单帧 PNG，保留透明像素。"""
    from PIL import Image, ImageOps
    # verify 检查文件结构，再重新打开解码；尺寸在分配像素缓冲区之前检查。
    with Image.open(source, formats=['PNG', 'JPEG', 'WEBP']) as image:
        if image.format.lower() != actual:
            raise ValueError('图片解码类型与文件头不符')
        if image.width * image.height > MAX_IMAGE_PIXELS or max(image.size) > 16384:
            raise ValueError('图片超过 1600 万像素或单边 16384 像素')
        if getattr(image, 'n_frames', 1) != 1:
            raise ValueError('暂不支持动图或多帧图片')
        image.verify()
    with Image.open(source, formats=['PNG', 'JPEG', 'WEBP']) as image:
        with ImageOps.exif_transpose(image) as oriented:
            original = oriented.size
            oriented.thumbnail((edge, edge), Image.Resampling.LANCZOS)
            return png_bytes(oriented), {'width':oriented.width, 'height':oriented.height,
                'original_width':original[0], 'original_height':original[1]}


def pdf_images(source, page_count, edge):
    # PDF 尺寸单位为点，比例 2 对应最高 144 DPI；较大页面按最长边降低比例。
    import pypdfium2 as pdfium
    images = {}
    with pdfium.PdfDocument(str(source)) as document:
        if len(document) != page_count:
            raise ValueError('PDF 页数校验不一致')
        for index in range(page_count):
            with closing(document[index]) as page:
                width, height = page.get_size()
                if not all(math.isfinite(n) and 0 < n <= 100000 for n in (width, height)):
                    raise ValueError('PDF 页面尺寸不合法')
                # 先计算比例再分配位图，避免巨型页面导致巨量内存分配。
                with closing(page.render(scale=min(2, edge / max(width, height)), may_draw_forms=False)) as bitmap:
                    with bitmap.to_pil() as image:
                        image.thumbnail((edge, edge))
                        images[f'page-{index + 1:03d}.png'] = png_bytes(image)
                if sum(map(len, images.values())) > MAX_TOTAL_BYTES:
                    raise ValueError('PDF 图片合计超过 4 MiB，请降低 max_long_edge')
    return images


def parse_file(task, workspace, out):
    """识别实际内容后解析；正文和可选图片统一接受产物大小限制。"""
    source = input_file(task, workspace)
    actual = detected_type(source)
    suffix = '.' + actual
    options = task.get('options', {})
    edge = options.get('max_long_edge', 2048)
    pages = options.get('max_pages', 7)
    if type(edge) is not int or not 64 <= edge <= 2048 or type(pages) is not int or not 1 <= pages <= 7:
        raise ValueError('解析参数超过限制')
    if options.get('pdf_images') and actual != 'pdf':
        raise ValueError('pdf_images 仅适用于 PDF')
    images = {}
    metadata = {}
    if suffix == ".zip":
        entries = check_zip(source)
        files = [item for item in entries if not item.is_dir()]
        if len(files) > 7:
            raise ValueError("ZIP 最多解出 7 个文件（另保留一份索引）")
        names = set()
        with zipfile.ZipFile(source) as archive:
            for index, entry in enumerate(files):
                # 输出重命名为扁平文件，原路径仅保留在索引中，彻底避免目录穿越。
                name = f"extracted-{index:02d}-{PurePosixPath(entry.filename).name}"
                if name in names:
                    raise ValueError("ZIP 包含重复文件名")
                names.add(name)
                with archive.open(entry) as stream:
                    data = stream.read(MAX_FILE_BYTES + 1)
                write_output(out, name, data)
        write_output_json(out, "parsed.json", {"source": source.name, "type": "zip", "files": [entry.filename for entry in files]})
        return
    if actual in ('png', 'jpeg', 'webp'):
        data, content = normalize_image(source, actual, edge)
        images['normalized.png'] = data
        content['image'] = 'normalized.png'
    elif suffix == ".csv":
        with source.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream, strict=True)
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > MAX_ROWS:
                    raise ValueError("CSV 超过 2000 行")
        content = rows
    elif suffix == ".json":
        content = json.loads(bounded_file(source).decode('utf-8-sig'))
    elif suffix == ".txt":
        content = bounded_file(source).decode("utf-8-sig")
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(source)
        if reader.is_encrypted:
            raise ValueError('暂不支持加密 PDF')
        if len(reader.pages) > 30:
            raise ValueError("PDF 超过 30 页")
        if options.get('pdf_images') and not 1 <= len(reader.pages) <= pages:
            raise ValueError(f'PDF 转图要求 1 至 {pages} 页，不会截断或丢页')
        content = "\n".join(page.extract_text() or "" for page in reader.pages)
        metadata['page_count'] = len(reader.pages)
        if options.get('pdf_images'):
            images = pdf_images(source, len(reader.pages), edge)
            metadata['images'] = list(images)
    elif suffix == ".docx":
        from docx import Document
        check_zip(source, 32 * 1024 * 1024)
        document = Document(source)
        content = {"paragraphs": [item.text for item in document.paragraphs],
                   "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]}
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        check_zip(source, 32 * 1024 * 1024)
        book = load_workbook(source, read_only=True, data_only=True)
        content = {}
        try:
            if len(book.worksheets) > 20:
                raise ValueError("Excel 超过 20 个工作表")
            for sheet in book.worksheets:
                if sheet.max_column > 200:
                    raise ValueError("Excel 工作表超过 200 列")
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) > MAX_ROWS:
                        raise ValueError("Excel 工作表超过 2000 行")
                content[sheet.title] = rows
        finally:
            book.close()
    else:
        raise ValueError("不支持的文件类型")
    if len(json.dumps(content, ensure_ascii=False, default=str)) > MAX_TEXT:
        raise ValueError("解析正文超过长度限制")
    report = (json.dumps({'source':source.name, 'type':actual, 'content':content, **metadata},
                        ensure_ascii=False, indent=2, default=str) + '\n').encode()
    # 所有产物先核验再写入，失败时不交回一部分页面冒充完整结果。
    if len(report) > MAX_FILE_BYTES or len(report) + sum(map(len, images.values())) > MAX_TOTAL_BYTES:
        raise ValueError('解析产物合计超过 4 MiB 或元数据超过 2 MiB')
    for name, data in images.items():
        write_output(out, name, data)
    write_output(out, 'parsed.json', report)


def audit_skill(task, workspace, out):
    # 只检查说明、引用和 Python AST，不执行被审核脚本，也不作全面安全认证。
    import yaml
    package = task_path(workspace, task.get("entrypoint", ""))
    if not package.is_dir() or package.is_symlink():
        raise ValueError("审核入口必须是技能包目录")
    issues = []
    def issue(file, message):
        issues.append({"file": file, "message": message})
    description = package / "SKILL.md"
    if not description.is_file():
        issue("SKILL.md", "缺少技能说明")
    else:
        text = bounded_file(description).decode("utf-8")
        match = re.match(r"\A---\s*\n(.*?)\n---(?:\s*\n|$)", text, re.S)
        try:
            fields = yaml.safe_load(match.group(1)) if match else {}
        except yaml.YAMLError:
            fields = {}
            issue("SKILL.md", "说明头部 YAML 无法解析")
        if not isinstance(fields, dict):
            fields = {}
        for name in ("name", "description"):
            if not isinstance(fields.get(name), str) or not fields[name].strip():
                issue("SKILL.md", f"缺少 {name} 字段")
        references = re.findall(r"\[[^\]]*\]\(([^)]+)\)", text)
        references += re.findall(r"`((?:scripts|references|assets)/[^`]+)`", text)
        for reference in sorted(set(references)):
            if reference.startswith(("http://", "https://", "#")):
                continue
            try:
                target = task_path(package, reference)
                if not target.exists() or target.is_symlink():
                    issue("SKILL.md", f"引用文件不存在：{reference}")
            except ValueError:
                issue("SKILL.md", f"引用路径越界：{reference}")
    scripts = sorted(package.rglob("*.py"))
    if len(scripts) > 50:
        raise ValueError("审核最多检查 50 个 Python 文件")
    suspicious = {"eval", "exec", "os.system", "subprocess.run", "subprocess.Popen"}
    for script in scripts:
        relative = str(script.relative_to(package))
        if script.is_symlink():
            issue(relative, "脚本是符号链接")
            continue
        try:
            tree = ast.parse(bounded_file(script).decode("utf-8"))
        except (SyntaxError, UnicodeError) as error:
            issue(relative, f"Python 语法错误：{error}")
            continue
        for node in ast.walk(tree):
            if isinstance(node, ast.Call):
                name = ast.unparse(node.func)
                if name in suspicious:
                    issue(relative, f"需要复核的调用：{name}")
    write_output_json(out, "audit.json", {"passed": not issues, "issues": issues,
                      "scope": "样例静态规则：说明字段、引用、Python 语法及有限可疑调用；不执行包，也不代表全面安全认证"})


def run_task(task, workspace: Path):
    # 此处生成的清单供管理侧参考，导出时仍独立枚举文件并重新计算摘要。
    kind = task.get("kind")
    if kind not in ("code", "skill", "web", "parse", "render", "audit"):
        raise ValueError("未知任务类型")
    if task.get("environment") not in ("python", "node", "browser", "files"):
        raise ValueError("未知环境类型")
    out = workspace / "out"
    out.mkdir(exist_ok=True)
    if out.is_symlink():
        raise ValueError("输出目录不能是符号链接")
    result = {"schema_version": 1, "kind": kind, "status": "success", "stdout": "", "stderr": "", "artifacts": []}
    try:
        if kind in ("code", "skill"):
            result.update(execute_script(task, workspace))
        elif kind == "web":
            browser_task(task, workspace, out)
        elif kind == "render":
            browser_task(task, workspace, out, render=True)
        elif kind == "parse":
            parse_file(task, workspace, out)
        elif kind == "audit":
            audit_skill(task, workspace, out)
        result["artifacts"] = [{key: item[key] for key in ("path", "size", "sha256")}
                               for item in read_artifacts(out, require_manifest=False)["artifacts"]]
    except Exception as error:
        result.update(getattr(error, "logs", {}))
        result["status"] = "failed"
        result["error"] = f"{type(error).__name__}: {error}"[:2000]
    write_output_json(out, "manifest.json", result)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--task", default="/workspace/task.json")
    args = parser.parse_args()
    try:
        task = json.loads(bounded_file(Path(args.task), 64 * 1024))
        result = run_task(task, Path("/workspace"))
        print(json.dumps({"status": result["status"], "kind": result["kind"]}, ensure_ascii=False))
        raise SystemExit(0 if result["status"] == "success" else 1)
    except Exception as error:
        print(f"[执行失败] {type(error).__name__}: {error}", file=sys.stderr)
        raise SystemExit(1)
