"""任务能力与产物边界测试；全部使用合成文件，不启动容器。"""
import base64
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from sandbox_model import guest
from sandbox_model.export import MAX_FILE_BYTES, read_artifacts

FIXTURES = Path(__file__).resolve().parents[1] / "tests" / "fixtures"


class GuestTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="guest-unit-")
        self.workspace = Path(self.temporary.name)
        (self.workspace / "inputs").mkdir()
        (self.workspace / "out").mkdir()

    def tearDown(self):
        self.temporary.cleanup()

    def run_task(self, kind, **options):
        return guest.run_task({"kind": kind, "environment": "python", **options}, self.workspace)

    def test_code_and_independent_export(self):
        (self.workspace / "main.py").write_text("from pathlib import Path\nPath('out/result.txt').write_text('合成结果')\nprint('任务完成')\n")
        result = self.run_task("code", entrypoint="main.py")
        self.assertEqual(result["status"], "success")
        exported = read_artifacts(self.workspace / "out")
        self.assertEqual(base64.b64decode(exported["artifacts"][0]["data_base64"]).decode(), "合成结果")
        self.assertIn("任务完成", exported["manifest"]["stdout"])

    def test_csv_parse(self):
        shutil.copy(FIXTURES / "input.csv", self.workspace / "inputs")
        result = self.run_task("parse")
        self.assertEqual(result["status"], "success")
        parsed = json.loads((self.workspace / "out/parsed.json").read_text())
        self.assertEqual(parsed["content"][0]["name"], "苹果")
        self.assertEqual(len(parsed["content"]), 2)

    def test_json_and_text_parse(self):
        source = self.workspace / "inputs/value.json"
        source.write_text('{"number":3}')
        result = self.run_task("parse", options={"input": "value.json"})
        self.assertEqual(result["status"], "success")
        self.assertEqual(json.loads((self.workspace / "out/parsed.json").read_text())["content"], {"number": 3})

    def test_docx_parse(self):
        from docx import Document
        document = Document()
        document.add_paragraph("合成 Word 输入")
        document.save(self.workspace / "inputs/document.docx")
        result = self.run_task("parse")
        self.assertEqual(result["status"], "success", result)
        parsed = json.loads((self.workspace / "out/parsed.json").read_text())
        self.assertEqual(parsed["content"]["paragraphs"], ["合成 Word 输入"])

    def test_xlsx_parse(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["商品", "数量"])
        workbook.active.append(["苹果", 2])
        workbook.save(self.workspace / "inputs/table.xlsx")
        result = self.run_task("parse")
        self.assertEqual(result["status"], "success", result)
        parsed = json.loads((self.workspace / "out/parsed.json").read_text())
        self.assertEqual(parsed["content"]["Sheet"][1], ["苹果", 2])

    def test_pdf_parse(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.write(self.workspace / "inputs/document.pdf")
        result = self.run_task("parse")
        self.assertEqual(result["status"], "success", result)
        parsed = json.loads((self.workspace / "out/parsed.json").read_text())
        self.assertEqual(parsed["type"], "pdf")

    def test_zip_extract_safe_names(self):
        with zipfile.ZipFile(self.workspace / "inputs/files.zip", "w") as archive:
            archive.writestr("folder/a.txt", "hello")
        result = self.run_task("parse")
        self.assertEqual(result["status"], "success")
        self.assertEqual((self.workspace / "out/extracted-00-a.txt").read_text(), "hello")

    def test_zip_traversal_fails(self):
        with zipfile.ZipFile(self.workspace / "inputs/files.zip", "w") as archive:
            archive.writestr("../escaped.txt", "unsafe")
        result = self.run_task("parse")
        self.assertEqual(result["status"], "failed")
        self.assertFalse((self.workspace / "escaped.txt").exists())

    def test_skill_audit_never_executes_package(self):
        for name, expected in (("skill_ok", True), ("skill_bad", False)):
            shutil.copytree(FIXTURES / name, self.workspace / name)
            result = self.run_task("audit", entrypoint=name)
            self.assertEqual(result["status"], "success")
            report = json.loads((self.workspace / "out/audit.json").read_text())
            self.assertEqual(report["passed"], expected)
            if not expected:
                self.assertGreaterEqual(len(report["issues"]), 3)
        self.assertFalse((self.workspace / "out/summary.json").exists())

    def test_log_overflow_fails(self):
        (self.workspace / "main.py").write_text("print('x' * 70000)\n")
        result = self.run_task("code", entrypoint="main.py")
        self.assertEqual(result["status"], "failed")
        self.assertIn("64 KiB", result["error"])

    def test_failed_script_keeps_bounded_logs(self):
        (self.workspace / "main.py").write_text("import sys\nprint('before failure')\nprint('synthetic failure', file=sys.stderr)\nsys.exit(3)\n")
        result = self.run_task("code", entrypoint="main.py")
        self.assertEqual(result["status"], "failed")
        self.assertIn("before failure", result["stdout"])
        self.assertIn("synthetic failure", result["stderr"])

    def test_entrypoint_traversal_fails(self):
        result = self.run_task("code", entrypoint="../outside.py")
        self.assertEqual(result["status"], "failed")

    def test_export_ignores_forged_manifest_paths(self):
        (self.workspace / "out/manifest.json").write_text('{"artifacts":[{"path":"/etc/passwd"}]}')
        (self.workspace / "out/actual.txt").write_text("safe")
        exported = read_artifacts(self.workspace / "out")
        self.assertEqual([item["path"] for item in exported["artifacts"]], ["actual.txt"])

    def test_export_refuses_symlink_and_hardlink(self):
        source = self.workspace / "inputs/input.txt"
        source.write_text("synthetic")
        link = self.workspace / "out/link.txt"
        link.symlink_to(source)
        with self.assertRaises(ValueError):
            read_artifacts(self.workspace / "out", require_manifest=False)
        link.unlink()
        os.link(source, link)
        with self.assertRaises(ValueError):
            read_artifacts(self.workspace / "out", require_manifest=False)

    def test_export_refuses_oversized_file(self):
        (self.workspace / "out/big.txt").write_bytes(b"x" * (MAX_FILE_BYTES + 1))
        with self.assertRaisesRegex(ValueError, "大小限制"):
            read_artifacts(self.workspace / "out", require_manifest=False)

    def test_export_refuses_many_empty_directories(self):
        for index in range(65):
            (self.workspace / "out" / str(index)).mkdir()
        with self.assertRaisesRegex(ValueError, "目录项"):
            read_artifacts(self.workspace / "out", require_manifest=False)

    def test_private_web_url_rejected(self):
        for url in ("http://127.0.0.1/", "http://169.254.169.254/", "http://user:pass@example.com/", "file:///etc/passwd"):
            with self.assertRaises(ValueError):
                guest.safe_url(url)
        self.assertEqual(guest.safe_url("https://example.com/"), "https://example.com/")


if __name__ == "__main__":
    unittest.main()
